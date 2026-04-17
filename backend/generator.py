import json
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
CONFIG_DIR = BASE_DIR / "config"
TEMPLATES_DIR = BASE_DIR.parent / "templates"


class CampaignConfigGenerator:
    def __init__(self):
        with open(CONFIG_DIR / "touchpoints.json") as f:
            self.touchpoints_config = json.load(f)
        with open(CONFIG_DIR / "market_defaults.json") as f:
            self.market_defaults = json.load(f)
        self.touchpoints = self.touchpoints_config["touchpoints"]

    def generate(
        self,
        campaign_name: str,
        campaign_type: str,
        markets: list,
        phases: list,
        global_promo_text: str,
        coupon_code: str = None,
    ) -> bytes:
        """Generate a campaign config spreadsheet and return it as bytes."""
        wb = Workbook()

        # --- Fills ---
        storyblok_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
        dy_fill = PatternFill(start_color="F3E5F5", end_color="F3E5F5", fill_type="solid")
        image_needed_fill = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
        header_fill = PatternFill(start_color="263238", end_color="263238", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)

        # Column layout
        columns = [
            ("Touchpoint", 25),
            ("System", 15),
            ("Phase", 15),
            ("Image URL", 50),
            ("Copy Text", 40),
            ("Link URL", 35),
            ("Start Date", 15),
            ("End Date", 15),
            ("Status", 12),
            ("Notes", 30),
        ]

        # --- Summary sheet ---
        ws_summary = wb.active
        ws_summary.title = "Summary"

        total_items = len(self.touchpoints) * len(phases) * len(markets)

        summary_data = [
            ("Campaign Name", campaign_name),
            ("Campaign Type", campaign_type.upper()),
            ("Markets", ", ".join(markets)),
            ("Phases", ", ".join(p["name"] for p in phases)),
            ("Date Range", f'{phases[0]["start_date"]} — {phases[-1]["end_date"]}'),
            ("Total Configuration Items", total_items),
            ("", ""),
            ("Completion Tracker", ""),
        ]

        ws_summary.column_dimensions["A"].width = 25
        ws_summary.column_dimensions["B"].width = 50

        title_font = Font(bold=True, size=12)
        for row_idx, (label, value) in enumerate(summary_data, start=1):
            cell_a = ws_summary.cell(row=row_idx, column=1, value=label)
            cell_a.font = title_font
            ws_summary.cell(row=row_idx, column=2, value=value)

        # Completion tracker per market
        tracker_row = len(summary_data) + 1
        for i, market in enumerate(markets):
            row = tracker_row + i
            ws_summary.cell(row=row, column=1, value=market).font = Font(bold=True)
            items_per_market = len(self.touchpoints) * len(phases)
            # COUNTIF formula referencing the market sheet's Status column (column I)
            formula = f"=COUNTIF('{market}'!I:I,\"Done\") & \" of {items_per_market} touchpoints configured\""
            ws_summary.cell(row=row, column=2, value=formula)

        # --- Market sheets ---
        for market in markets:
            ws = wb.create_sheet(title=market)
            defaults = self.market_defaults.get(market, {})

            # Write headers
            for col_idx, (col_name, col_width) in enumerate(columns, start=1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
                ws.column_dimensions[get_column_letter(col_idx)].width = col_width

            # Freeze header row
            ws.freeze_panes = "A2"

            # Auto-filters
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

            # Write data rows
            row_num = 2
            for phase in phases:
                for tp in self.touchpoints:
                    # Build copy text
                    copy_text = ""
                    if "copy_text" in tp["fields"]:
                        parts = [p for p in [global_promo_text, defaults.get("promo_suffix", "")] if p]
                        copy_text = " | ".join(parts)

                    # Link URL
                    link_url = ""
                    if "link_url" in tp["fields"]:
                        link_url = defaults.get("sale_page_url", "")

                    # Notes
                    notes = ""
                    if tp["name"] == "Notification Bar":
                        notes = f'background_color: {defaults.get("notification_bg", "")}'
                    if tp["name"] == "Promotional Pop-up" and coupon_code:
                        notes = f"coupon_code: {coupon_code}"

                    # Image URL
                    image_url = ""

                    ws.cell(row=row_num, column=1, value=tp["name"])
                    ws.cell(row=row_num, column=2, value=tp["system"])
                    ws.cell(row=row_num, column=3, value=phase["name"])
                    img_cell = ws.cell(row=row_num, column=4, value=image_url)
                    ws.cell(row=row_num, column=5, value=copy_text)
                    ws.cell(row=row_num, column=6, value=link_url)
                    ws.cell(row=row_num, column=7, value=phase["start_date"])
                    ws.cell(row=row_num, column=8, value=phase["end_date"])
                    ws.cell(row=row_num, column=9, value="Not Started")
                    ws.cell(row=row_num, column=10, value=notes)

                    # Row fill based on system
                    row_fill = storyblok_fill if tp["system"] == "Storyblok" else dy_fill
                    for col_idx in range(1, len(columns) + 1):
                        ws.cell(row=row_num, column=col_idx).fill = row_fill

                    # Yellow fill for Image URL if needs_image
                    if tp["needs_image"]:
                        img_cell.fill = image_needed_fill

                    row_num += 1

            # Conditional formatting for Status column (column I)
            last_row = row_num - 1
            if last_row >= 2:
                status_range = f"I2:I{last_row}"
                ws.conditional_formatting.add(
                    status_range,
                    CellIsRule(
                        operator="equal",
                        formula=['"Done"'],
                        fill=PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
                    ),
                )
                ws.conditional_formatting.add(
                    status_range,
                    CellIsRule(
                        operator="equal",
                        formula=['"Blocked"'],
                        fill=PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid"),
                    ),
                )

        # Return as bytes
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def save_template(self, campaign_name: str, config: dict) -> str:
        """Save a campaign config as a reusable JSON template."""
        TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
        date_str = datetime.now().strftime("%Y%m%d")
        safe_name = "".join(c if c.isalnum() or c in "-_" else "_" for c in campaign_name).strip("_")
        filename = f"{safe_name}_{date_str}.json"
        filepath = TEMPLATES_DIR / filename
        with open(filepath, "w") as f:
            json.dump(config, f, indent=2)
        return filename

    def list_templates(self) -> list:
        """Return all saved template filenames."""
        TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
        return sorted(f.name for f in TEMPLATES_DIR.glob("*.json"))

    def load_template(self, filename: str) -> dict:
        """Read and return a saved template."""
        filepath = (TEMPLATES_DIR / filename).resolve()
        if not str(filepath).startswith(str(TEMPLATES_DIR.resolve())):
            raise FileNotFoundError("Invalid filename")
        with open(filepath) as f:
            return json.load(f)
